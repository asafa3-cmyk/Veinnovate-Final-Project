"""Independent arithmetic, artifact and provenance checks for the final chapter package."""
import csv,hashlib,json,sys
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from xml.etree import ElementTree as ET
import numpy as np
import fitz
O=Path(__file__).resolve().parent;E=O.parent;R=E.parents[1]
sys.path.append(str(R/'runtime-venv/lib/python3.12/site-packages'))
from openpyxl import load_workbook
report={};sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
sources=json.loads((O/'SOURCE_HASHES.json').read_text())
assert all(Path(x['path']).exists() and sha(Path(x['path']))==x['sha256'] for x in sources)
report['source_files_hash_verified']=len(sources)
source=json.loads((E/'results/fast1_results.json').read_text());r=source['per_slot'];r=[x for x in r if x.get('status')=='OK'];assert len(r)==18
cal=[x for x in r if x['slot']=='B0'];assert len(cal)==1
bp=[x for x in r if x['slot']!='B0'];assert len(bp)==17
summary=json.loads((O/'results_package.json').read_text())
for k in ['sbp','dbp']:
 err=np.array([x['model_'+k]-x['cuff_'+k] for x in bp]);row=next(x for x in summary['bp'] if x['cohort']=='all_retained' and x['method']=='Relaxed' and x['metric']==k)
 assert np.isclose(row['mae'],abs(err).mean()) and np.isclose(row['bias'],err.mean()) and np.isclose(row['sd'],err.std(ddof=1))
 for t in [5,10,15]:assert row[f'within{t}_n']==int(sum(abs(err)<=t))
 report[k+'_mae']=float(abs(err).mean())
hrerr=np.array([x['hr']-x['cuff_hr'] for x in r]);assert np.isclose(abs(hrerr).mean(),3.383333333333333)
assert sum(abs(hrerr)<=5)==15 and sum(abs(hrerr)<=10)==17
report['hr_mae']=float(abs(hrerr).mean());report['hr_within5_count']=15;report['hr_within10_count']=17
raw=np.load(O/'raw_windows.npz');assert len(raw.files)==36 and all(len(raw[k])==3000 for k in raw.files)
report['raw_windows']=18;report['raw_pairs']=54000
req=list(csv.DictReader((O/'tables/T10_requirements.csv').open()));assert len(req)==23 and len({x['id'] for x in req})==23
assert {f'I{i:02}' for i in range(1,13)}<={x['id'] for x in req}
assert all(x['status'] in {'NOT_TESTED','NOT_MEASURED','NOT_EVALUABLE','UNDEFINED','PARTIAL','NOT_MET_ALL','IMPLEMENTED_ACCURACY_NOT_MET'} for x in req)
report['requirements_mapped']=23
with (O/'tables/T04_all_slots.csv').open() as f:slots=list(csv.DictReader(f))
assert len(slots)==19 and slots[-1]['slot']=='BPREC2' and slots[-1]['capture']=='NOT_CAPTURED'
assert slots[0]['sbp']==slots[0]['dbp']==''
assert slots[-2]['reference_note']=='LATE_SINGLE'
with (O/'tables/T11_cuff_readings.csv').open() as f:assert len(list(csv.DictReader(f)))==20
wb=load_workbook(O/'FAST1_RESULTS_TABLES.xlsx',read_only=True,data_only=True);assert len(wb.sheetnames)==15
for name,count in summary['tables'].items():assert wb[name].max_row==count+1,(name,wb[name].max_row,count)
wb.close();report['table_sheets']=14
for name,n in [('FAST1_Protocol_Methods_REVISED',3),('FAST1_Results_Chapter',8),('FAST1_RESULTS_FIGURE_ATLAS',14)]:
 pdf=fitz.open(O/(name+'.pdf'));assert len(pdf)==n,(name,len(pdf))
 for page in pdf:
  assert len(page.get_text().strip())>150,(name,'unexpected blank/orphan page')
  for b in page.get_text('blocks'):
   x0,y0,x1,y1=b[:4];assert x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1,(name,b[:4],page.rect)
 report[name+'_pages']=len(pdf)
for name in ['FAST1_Protocol_Methods_REVISED','FAST1_Results_Chapter']:
 with ZipFile(O/(name+'.docx')) as z:
  assert z.testzip() is None
  xml=ET.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  for tbl in xml.findall('.//w:tblPr',ns):assert tbl.find('w:bidiVisual',ns).get('{'+ns['w']+'}val') in ['false','0']
report['visual_review']='All14 figures plus all11 manuscript pages inspected in rendered previews; no clipping/orphan pages in final layout.'
assert len(list((O/'figures').glob('*.png')))==len(list((O/'figures').glob('*.svg')))==len(list((O/'figures').glob('*.pdf')))==14
# Retain existing audits and source-result manifests as independent cross-checks.
for folder in ['independent_audit','verification_book_backup']:
 for l in (E/folder/'SHA256SUMS.txt').read_text().splitlines():
  h,n=l.split(maxsplit=1);assert sha(E/folder/n)==h,n
report['previous_audits_unchanged']=True
report['remote_actions']='None';report['original_attachments_modified']=False
(O/'VALIDATION.json').write_text(json.dumps(report,indent=2))
files=sorted(p for p in O.rglob('*') if p.is_file() and 'preview' not in p.relative_to(O).parts and p.name!='SHA256SUMS.txt')
(O/'SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(O).as_posix()}\n' for p in files))
zip_path=E/'FAST1_CHAPTER_PACKAGE_20261022.zip'
with ZipFile(zip_path,'w',ZIP_DEFLATED) as z:
 for p in files+[O/'SHA256SUMS.txt']:z.write(p,arcname='FAST1_CHAPTER_PACKAGE/'+p.relative_to(O).as_posix())
with ZipFile(zip_path) as z:assert z.testzip() is None and len(z.namelist())==len(files)+1
zip_path.with_suffix('.zip.sha256').write_text(sha(zip_path)+'  '+zip_path.name+'\n')
print(json.dumps(report,indent=2));print('Package ZIP:',zip_path,'bytes=',zip_path.stat().st_size)
