"""Generate reproducible FAST1 results tables and publication figures from verified retained data."""
import os
os.environ['MPLCONFIGDIR']='/tmp/fast1_mpl_chapter'
import csv,json,hashlib,math
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import numpy as np
from scipy.stats import spearmanr
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.ticker import MaxNLocator
import sys
sys.path.append(str(Path(__file__).resolve().parents[3]/"runtime-venv/lib/python3.12/site-packages"))
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
O=Path(__file__).resolve().parent;E=O.parent;R=E.parents[1];RUN=Path.home()/'VeinnovateRuns/POC-20261021-FAST1'
F=O/'figures';T=O/'tables';F.mkdir(exist_ok=True);T.mkdir(exist_ok=True)
J=json.loads((E/'results/fast1_results.json').read_text());P=J['proto'];S=J['strict'];plan=json.loads((E/'fast1_plan.json').read_text());old=json.loads((E/'independent_audit/audit.json').read_text())
slots=[s['slot'] for s in plan['slots']];kinds={s['slot']:s['kind'] for s in plan['slots']};captured=[s for s in slots if s in P];test=[s for s in captured if s!='B0'];sealed={e['slot']:e for e in map(json.loads,(RUN/'jetson_sealed_records.jsonl').read_text().splitlines())};live={e['slot']:e for e in map(json.loads,(RUN/'live_results.jsonl').read_text().splitlines())}
events=list(map(json.loads,(RUN/'cuff_events.jsonl').read_text().splitlines()));saves={e['slot']:e for e in events if e['kind']=='SAVE'}
C={s:dict(zip(['sbp','dbp','hr'],np.mean(e.get('readings') or [e['values']],axis=0))) for s,e in saves.items()}
def ts(s):return datetime.fromisoformat(s.replace('Z','+00:00'))
def local(s):return ts(s).astimezone(ZoneInfo('Asia/Jerusalem')).strftime('%H:%M:%S')
def num(x):return x is not None and np.isfinite(x)
def stats(pred,ref,eligible):
 e=np.asarray(pred,float)-np.asarray(ref,float);n=len(e)
 if n==0:return dict(n=0,eligible=eligible,coverage_pct=0)
 sd=float(e.std(ddof=1)) if n>1 else None
 return dict(n=n,eligible=eligible,coverage_pct=100*n/eligible,mae=float(np.mean(abs(e))),bias=float(e.mean()),sd=sd,rmse=float(np.sqrt(np.mean(e**2))),within5_n=int(sum(abs(e)<=5)),within10_n=int(sum(abs(e)<=10)),within15_n=int(sum(abs(e)<=15)),within5_pct=100*float(np.mean(abs(e)<=5)),within10_pct=100*float(np.mean(abs(e)<=10)),within15_pct=100*float(np.mean(abs(e)<=15)),loa_low=float(e.mean()-1.96*sd) if sd is not None else None,loa_high=float(e.mean()+1.96*sd) if sd is not None else None)
assert stats([5,15],[10,10],3)['mae']==5
start=ts(sealed['B0']['ppg_start_utc']);minutes={s:(ts(sealed[s]['ppg_start_utc'])-start).total_seconds()/60 for s in captured}
cohorts={'all_retained':test,'exclude_late_BPUP2':[s for s in test if s!='BPUP2'],'exclude_late_and_BPUP1_disagreement':[s for s in test if s not in ['BPUP1','BPUP2']]}
bp=[]
for cohort,ss in cohorts.items():
 strict=[s for s in ss if S[s]['sbp'] is not None]
 for method,chosen in [('Relaxed',ss),('Strict',strict),('B0 constant',ss),('B0 on strict subset',strict)]:
  for metric in ['sbp','dbp']:
   pred=[C['B0'][metric] if method.startswith('B0') else (S if method=='Strict' else P)[s][metric] for s in chosen]
   bp.append(dict(cohort=cohort,method=method,metric=metric,**stats(pred,[C[s][metric] for s in chosen],len(ss))))
hr=[]
for cohort,ss in [('all_retained',captured),('exclude_late_BPUP2',[s for s in captured if s!='BPUP2']),('nonbaseline',[s for s in captured if s!='B0'])]:
 for tier in ['ALL','HIGH','ARTIFACT']:
  ch=[s for s in ss if tier=='ALL' or P[s]['hr_tier']==tier];st=stats([P[s]['hr'] for s in ch],[C[s]['hr'] for s in ch],len(ss))
  st['mape_pct']=float(np.mean([abs(P[s]['hr']-C[s]['hr'])/C[s]['hr'] for s in ch])*100)
  hr.append(dict(cohort=cohort,tier=tier,**st))
per=[];timing=[];cuff=[]
for s in slots:
 if s not in P:per.append(dict(slot=s,condition=kinds[s],capture='NOT_CAPTURED'));continue
 q=sealed[s];l=live.get(s,{});rd=saves[s].get('readings') or [saves[s]['values']];diff=max(r[0] for r in rd)-min(r[0] for r in rd)
 per.append(dict(slot=s,condition=kinds[s],capture='COMPLETE',start_local=local(q['ppg_start_utc']),end_local=local(q['ppg_end_utc']),start_utc=q['ppg_start_utc'],minutes_from_B0=minutes[s],cuff_sbp=C[s]['sbp'],cuff_dbp=C[s]['dbp'],cuff_hr=C[s]['hr'],cuff_n=len(rd),reference_note='LATE_SINGLE' if s=='BPUP2' else 'REPEAT_SBP_DISAGREEMENT_16' if s=='BPUP1' else 'OVERLAP_NOT_INDEPENDENTLY_TIMED',sbp=q['sbp'],dbp=q['dbp'],sbp_error=None if s=='B0' else q['sbp']-C[s]['sbp'],dbp_error=None if s=='B0' else q['dbp']-C[s]['dbp'],hr=P[s]['hr'],hr_error=P[s]['hr']-C[s]['hr'],hr_tier=P[s]['hr_tier'],spo2_experimental=P[s]['spo2'],strict_state=S[s]['bp_state'],strict_sbp=None if s=='B0' else S[s]['sbp'],strict_dbp=None if s=='B0' else S[s]['dbp'],warnings=P[s]['warnings'],pairs=q['measurement_pairs'],packet_loss=q['packet_loss'],live_latency_s=l.get('latency_s')))
 intent=next((e['t'] for e in events if e['kind']=='CAPTURE_INTENT' and e['slot']==s),None)
 dur=(ts(q['ppg_end_utc'])-ts(q['ppg_start_utc'])).total_seconds()
 timing.append(dict(slot=s,request_mac_utc=intent,start_edge_utc=q['ppg_start_utc'],end_edge_utc=q['ppg_end_utc'],sealed_span_s=dur,live_acquisition_span_s=(ts(l['acq_end'])-ts(l['acq_start'])).total_seconds() if l else None,live_post_acquisition_s=l.get('latency_s'),emitted_utc=l.get('emitted'),save_mac_utc=saves[s]['t'],note='Mac request/SAVE and edge acquisition are different clocks; no cuff measurement interval is recorded'))
 for i,r in enumerate(rd):cuff.append(dict(slot=s,reading_index=i+1,sbp=r[0],dbp=r[1],pulse=r[2],save_utc=saves[s]['t'],is_late=s=='BPUP2',repeat_disagreement=diff>10))
coverage=[dict(measure='Captured windows',numerator=18,denominator=19,pct=100*18/19,meaning='Planned slots, including calibration'),dict(measure='BP on captured slots',numerator=17,denominator=18,pct=100*17/18,meaning='B0 has no live BP'),dict(measure='BP on planned slots',numerator=17,denominator=19,pct=100*17/19,meaning='B0 calibration plus missing BPREC2'),dict(measure='BP on captured evaluation slots',numerator=17,denominator=17,pct=100,meaning='Excludes B0'),dict(measure='BP on planned evaluation slots',numerator=17,denominator=18,pct=100*17/18,meaning='Excludes B0, includes missing BPREC2'),dict(measure='Strict BP evaluation coverage',numerator=11,denominator=17,pct=100*11/17,meaning='Offline quality gate; 6 withheld'),dict(measure='HR numerical output',numerator=18,denominator=19,pct=100*18/19,meaning='Includes B0; planned slots'),dict(measure='HR HIGH tier',numerator=14,denominator=18,pct=100*14/18,meaning='Captured slots'),dict(measure='SpO2 numerical output',numerator=18,denominator=19,pct=100*18/19,meaning='Uncalibrated; no accuracy evidence'),dict(measure='Live latency records',numerator=17,denominator=18,pct=100*17/18,meaning='BPUP2 missing from driver live log')]
req=[]
def rq(id,section,requirement,target,observed,status,exhibit,limit):req.append(dict(id=id,attachment_section=section,requirement=requirement,target=target,observed=observed,status=status,exhibit=exhibit,limit=limit))
rq('G1','2.2.1 p1','Accurate absolute BP; continuous; without cuff','Absolute BP from PPG','17 numerical BP estimates; 18 periodic 30-s windows; initial cuff calibration','PARTIAL','F01,F02,T01,T02','Not calibration-free or beat-to-beat; full-cohort MAE fails')
rq('G2','2.2.1 p1','Oxygen saturation measurement','Reference accuracy required','18 uncalibrated numeric estimates, no oximeter','NOT_EVALUABLE','F11,T09','Numeric availability is not accuracy')
rq('G3','2.2.1 p1','Real-time endpoint monitoring','Reliable transmission and display','17 edge emission latency records','PARTIAL','F10,T07','No dashboard-render measurement')
rq('G4','2.2.1 p1','Integration with intravenous cannula','Long-term integration vision','No mass/mechanical test in FAST1','NOT_TESTED','T10','Do not infer mechanical validation from dorsal placement')
rq('O1','2.2.2 p1','ML/DL absolute BP from raw pulse','Frozen model inference','Actual sensor-to-model predictions reproduced','IMPLEMENTED_ACCURACY_NOT_MET','F02,T02','No new training; trained model uses baseline cuff')
rq('O2','2.2.2 p1','High-quality HR and oxygen saturation','Independent accuracy evidence','HR MAE 3.38 BPM; oxygen reference absent','PARTIAL','F07,F08,F11,T05','HR scalar is not a separate model input; shared waveform quality matters')
rq('O3','2.2.2 p1','Big Data and diverse-population generalization','Training/held-out dataset evidence','One adult session only','NOT_TESTED','T10','Historical dataset claims not revalidated by FAST1')
rq('O4','2.2.2 p1','Clinical interface and dynamic environment','Workflow and variable-condition testing','Rest/motion/challenge slots; no nurse-display trial','PARTIAL','F01,F09,T06','Not clinical workflow validation')
rq('O5','2.2.2 p1-2','Reliable real-time transport and noise reduction','Retained raw, processing and dedicated display','274 hashed chunks, 18 complete windows; gate comparison','PARTIAL','F06,F10,F12a,F12b,T08','Within-window integrity does not prove entire-session continuity')
rq('O6','2.2.2 p2','Full component synchronization','Temporal alignment','Slot association recorded; late BPUP2; no cuff start/end','PARTIAL','T04,T07','Slot ID is not proof of temporal overlap')
rq('O7','2.2.2 p2','Prior prediction accuracy exceeding 75%','Definition and prior validation provenance','Accuracy denominator/tolerance unspecified','UNDEFINED','T10','Do not use 100-MAPE, AUC, or fraction of valid results as substitute')
for id,k,targ,obs in [('I01','SBP','MAE 8-10 mmHg','13.36 all; 12.48 excluding late BPUP2; 9.29 strict 11/17'),('I02','DBP','MAE 5-7 mmHg','7.87 all; 7.34 excluding late BPUP2; 6.27 strict 11/17')]:rq(id,'2.2.3 p2',k+' MAE',targ,obs,'NOT_MET_ALL','F02,F03,F06,T02','Upper limits 10/7; lower 8/5 are stronger targets; lower errors are better')
rq('I03','2.2.3 p2','HR accuracy, typical and dynamic','MAE 6-10 BPM; >=85% agreement','MAE3.38; within5 83.3%; within10 94.4%','PARTIAL','F07,F08,T05','Agreement tolerance absent; cuff pulse is not independently verified ECG. No unconditional agreement PASS')
rq('I04','2.2.3 p2','SpO2 accuracy','Most within3 points; >=90% within4 points','No reference readings','NOT_EVALUABLE','F11,T09','Reference cuff recorded BP/pulse only')
rq('I05','2.2.3 p2','BP distribution (BHS-style)','>=40/65/85% within5/10/15 mmHg','SBP17.6/47.1/52.9%; DBP41.2/64.7/88.2%','NOT_MET_ALL','F05,T03','Descriptive distribution, not a formal BHS grade; 64.7% must not round to a pass at65%')
rq('I06','2.2.3 p2','Processing and display time','30-120 seconds','30-s nominal acquisition; live post-acquisition p95 0.274s','PARTIAL','F10,T07','No command-to-render measurement; subsecond latency alone cannot pass this requirement')
rq('I07','2.2.3 p2','Transmission range','4m; up to8m optimal','Distance not measured','NOT_TESTED','T10','No invented range plot')
rq('I08','2.2.3 p2','Dedicated display','Smooth presentation','No dashboard trial','NOT_TESTED','T10','No fabricated usability or render latency')
rq('I09','2.2.3 p2','Full system synchronization','Continuous efficient operation','18/19 windows; timestamps; cuff overlap uncertain','PARTIAL','F01,T01,T07','Empty event log does not establish zero disconnects')
rq('I10','2.2.3 p2','Mounted mass excluding batteries','<= mass of3 equivalent cannulas','No weighing evidence','NOT_MEASURED','T10','No fabricated mass comparison')
rq('I11','2.2.3 p2','Prior BP accuracy >75%','Undefined regression accuracy','Not derivable from MAE without definition','UNDEFINED','T10','Duplicates historical assertion O7')
rq('I12','2.2.3 p2','Broader age range including children','Population applicability','One adult only','NOT_TESTED','T10','No pediatric inference')
assert len(req)==23
signal=[];raw=np.load(O/'raw_windows.npz')
for s in captured:
 for ch in ['red','ir']:
  x=raw[s+'_'+ch];signal.append(dict(slot=s,channel=ch,n=len(x),median_counts=float(np.median(x)),min_counts=float(x.min()),max_counts=float(x.max()),sd_counts=float(x.std()),preflight_status=sealed[s]['preflight_status'],preflight_adjustments=sealed[s]['preflight_adjustments']))
sp=[dict(slot=s,value=P[s]['spo2'],label='UNCALIBRATED_MAXIM_LINEAR',reference=None,accuracy='NOT_EVALUABLE') for s in captured]
tables={'T01_coverage':coverage,'T02_BP_summary':bp,'T03_error_distribution':[r for r in bp if r['cohort']=='all_retained'],'T04_all_slots':per,'T05_HR_summary':hr,'T06_change_tracking':J['metrics']['change_tracking'],'T07_timing':timing,'T08_integrity':old['tables']['Integrity'],'T09_SpO2':sp,'T10_requirements':req,'T11_cuff_readings':cuff,'T12_HR_algorithms':old['tables']['HR_comparison'],'T13_HR_per_slot':old['tables']['HR_per_slot'],'T14_signal_summary':signal}
wb=Workbook();wb.remove(wb.active);ws=wb.create_sheet('START_HERE')
notes=[('Source','Two-page requirements screenshot supplied 2026-10-22; thresholds refer to that attachment, not additional targets from other book drafts.'),('Evaluation','All retained 18 windows are shown; BP excludes calibration B0 (n=17). BPREC2 is missing. Late BPUP2 is flagged and excluded in sensitivity analysis.'),('Gates','Relaxed selected quality gates are disabled; integrity/model/baseline gates remain. Strict is offline replay. Shared HR HIGH mask is exploratory, not a validated BP gate.'),('Statistics','MAE, signed bias, sample SD and descriptive limits of agreement. Same-person repeated windows are not independent subjects; no population CI or formal validation claim.'),('References','B0 mean126/74,pulse90. BPUP1 repeat SBP differs16mmHg. BPUP2 late single. SAVE timestamps are entry times, not cuff acquisition intervals.'),('Coverage','Report denominators. B0 has no live BP; replay B0 self-prediction is not used. Missing BPREC2 remains visible.'),('Unmeasured','No oxygen reference, range/mass/clinical UI/population test. No fake plots for unavailable measurements.'),('Versions','prototype-continuous-1.1.0; ONNX1f017653; firmware sealed3.0.5; sensor100Hz effective, 3000 pairs/window.'),('Reuse','All originals are preserved. Captions and requirements mapping identify limitations; figures exported PNG300dpi, SVG and PDF.')]
for r in notes:ws.append(r)
ws.column_dimensions['A'].width=22;ws.column_dimensions['B'].width=110
for row in ws:
 row[1].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=50
for name,rows in tables.items():
 cols=list(dict.fromkeys(k for r in rows for k in r));ws=wb.create_sheet(name);ws.append(cols)
 for row in rows:ws.append([json.dumps(row.get(c)) if isinstance(row.get(c),(list,dict)) else row.get(c) for c in cols])
 ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
 for col in ws.columns:
  ws.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col)+1))
  for c in col:
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,float):c.number_format='0.00'
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17354D')
 with (T/(name+'.csv')).open('w') as f:w=csv.DictWriter(f,fieldnames=cols);w.writeheader();w.writerows(rows)
wb.save(O/'FAST1_RESULTS_TABLES.xlsx')
# Publication palette and consistent styling.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.grid':True,'grid.alpha':.18,'figure.facecolor':'white','axes.titleweight':'bold','savefig.facecolor':'white'})
blue='#23699B';orange='#D37824';green='#398266';red='#B64949';gray='#78838C';colors={'CAL':gray,'REST':blue,'MOTION':orange,'BPUP':red,'BPREC':green}
figures=[]
from textwrap import fill
atlas=PdfPages(O/'FAST1_RESULTS_FIGURE_ATLAS.pdf')
def finish(fig,id,title,caption):
 fig.suptitle(title,fontsize=15,fontweight='bold',x=.05,ha='left')
 fig.text(.05,.02,fill(caption,150),fontsize=8,va='bottom',color='#35454F')
 fig.tight_layout(rect=(.02,.12,.98,.92))
 for ext in ['png','svg','pdf']:fig.savefig(F/(id+'.'+ext),dpi=300)
 atlas.savefig(fig);plt.close(fig);figures.append(dict(id=id,title=title,caption=caption,file='figures/'+id+'.png'))
def points(ax,ss,x,y):
 for s,xx,yy in zip(ss,x,y):ax.scatter(xx,yy,c=colors[kinds[s]],s=45,edgecolors='black' if s=='BPUP2' else 'white',linewidths=1.4 if s=='BPUP2' else .5,marker='D' if s=='BPUP2' else 'o',zorder=3)
def slot_axis(ax,ss):ax.set_xticks(range(len(ss)),ss,rotation=55,ha='right');ax.set_xlabel('Chronological slot (unequal real-time spacing)')
fig,axs=plt.subplots(1,2,figsize=(13,7),gridspec_kw={'width_ratios':[1.7,1]})
ax=axs[0]
for i,s in enumerate(slots):
 if s in sealed:ax.barh(i,(ts(sealed[s]['ppg_end_utc'])-ts(sealed[s]['ppg_start_utc'])).total_seconds()/60,left=minutes[s],color=colors[kinds[s]],height=.65);ax.text(minutes[s]+.7,i,local(sealed[s]['ppg_start_utc']),fontsize=8,va='center')
 else:ax.text(1,i,'NOT CAPTURED',color=red,va='center')
ax.set_yticks(range(len(slots)),slots);ax.invert_yaxis();ax.set_xlim(-1,62);ax.set_xlabel('Minutes from first B0 notification');ax.set_title('18 captured windows across 55.8 min')
cov=coverage[:3]+[coverage[5],coverage[6],coverage[9]];ys=np.arange(len(cov));axs[1].barh(ys,[r['pct'] for r in cov],color=[blue,blue,blue,orange,green,gray]);axs[1].set_yticks(ys,[r['measure'].replace(' ','\n',1) for r in cov],fontsize=8);axs[1].invert_yaxis();axs[1].set_xlim(0,125);axs[1].set_xlabel('Coverage (%)')
for i,r in enumerate(cov):axs[1].text(r['pct']+1,i,f"{r['numerator']}/{r['denominator']}",va='center')
finish(fig,'F01_timeline_coverage','F01 | What was actually recorded','Thirty-second windows were scheduled at roughly three-minute start-to-start intervals; challenge slots required manual GO. Total nominal sampled time is 9 min. Blank intervals are unsampled scheduled gaps, not demonstrated packet loss. BPREC2 is missing; B0 is calibration with no live BP.')
fig,axs=plt.subplots(2,1,figsize=(13,8),sharex=True)
for ax,k in zip(axs,['sbp','dbp']):
 x=[minutes[s] for s in test];ax.plot(x,[C[s][k] for s in test],'-o',color=blue,label='Cuff reference');ax.plot(x,[P[s][k] for s in test],'-o',color=orange,label='VeinNet relaxed');ax.axhline(C['B0'][k],color=gray,ls='--',label='B0 constant')
 ax.scatter(minutes['BPUP2'],P['BPUP2'][k],marker='D',s=85,facecolors='none',edgecolors='black',zorder=5);ax.set_ylabel(k.upper()+' (mmHg)');ax.legend(ncol=3,fontsize=9)
 for s in ['MOT1','BPUP1','BPREC1','MOT2','BPUP2']:ax.axvline(minutes[s],color=colors[kinds[s]],alpha=.25)
axs[1].set_xticks([minutes[s] for s in test],test,rotation=55,ha='right');axs[1].set_xlabel('Slot at actual elapsed time from B0')
finish(fig,'F02_BP_trajectory','F02 | BP estimates versus cuff over the session','B0 excluded from evaluation; n=17. Lines connect intermittent observations, not continuous BP measurements. BPUP1 uses a disagreeing cuff pair; BPUP2 (open diamond) uses a late single reading. Separate reference-quality sensitivity tables are provided.')
fig,axs=plt.subplots(1,2,figsize=(11,6))
for ax,k in zip(axs,['sbp','dbp']):
 x=[C[s][k] for s in test];y=[P[s][k] for s in test];lo=min(x+y)-3;hi=max(x+y)+3;ax.plot([lo,hi],[lo,hi],'--',color=gray,label='Identity');points(ax,test,x,y);ax.set(xlim=(lo,hi),ylim=(lo,hi),xlabel='Cuff '+k.upper()+' (mmHg)',ylabel='Predicted '+k.upper()+' (mmHg)');ax.set_aspect('equal');ax.set_title(k.upper()+' | n=17')
finish(fig,'F03_BP_agreement','F03 | Paired BP agreement','The diagonal is perfect agreement. Colors: rest blue, motion orange, challenge red, recovery green. Late BPUP2 is a black-edged diamond. B0 is excluded; no regression fit or post-hoc offset is applied.')
fig,axs=plt.subplots(1,3,figsize=(14,6))
for ax,k,ss in zip(axs,['sbp','dbp','hr'],[test,test,captured]):
 x=np.array([(P[s][k]+C[s][k])/2 for s in ss]);y=np.array([P[s][k]-C[s][k] for s in ss]);st=stats([P[s][k] for s in ss],[C[s][k] for s in ss],len(ss));points(ax,ss,x,y)
 for v,label,col,style in [(st['bias'],'Bias',blue,'-'),(st['loa_low'],'Lower LoA',red,'--'),(st['loa_high'],'Upper LoA',red,'--')]:ax.axhline(v,color=col,ls=style,label=f'{label}: {v:.2f}')
 ax.axhline(0,color=gray,alpha=.5);ax.set_title(f'{k.upper()} | n={len(ss)}');ax.set_xlabel('Mean of estimate and reference');ax.set_ylabel('Estimate - reference');ax.legend(fontsize=8,loc='lower left')
finish(fig,'F04_Bland_Altman','F04 | Descriptive bias and limits of agreement','Bias +/-1.96 sample SD; BP units mmHg, HR units BPM. Limits describe these repeated windows from one participant and are not population limits or clinical validation. BPUP2 late reference retained and marked; BP baseline excluded.')
fig,axs=plt.subplots(1,2,figsize=(12,6))
for ax,k in zip(axs,['sbp','dbp']):
 for method,data,ss,col in [('Relaxed',P,test,blue),('Strict',S,[s for s in test if S[s][k] is not None],orange)]:
  errs=np.sort([abs(data[s][k]-C[s][k]) for s in ss]);ax.step(np.r_[0,errs],np.r_[0,np.arange(1,len(errs)+1)/len(errs)*100],where='post',label=f'{method} n={len(ss)}',color=col)
 ax.scatter([5,10,15],[40,65,85],marker='X',color=red,s=70,label='Attachment target',zorder=4);ax.set(xlim=(0,35),ylim=(0,105),xlabel='Absolute error (mmHg)',ylabel='Within threshold (%)');ax.set_title(k.upper());ax.legend(fontsize=9)
finish(fig,'F05_error_distribution','F05 | Error distributions against the exact attached targets','Required cumulative fractions are >=40%, >=65%, >=85% within 5,10,15 mmHg. Full-cohort DBP reaches 64.7%, not65%, within10mmHg. Strict curves have a different denominator (11 vs17); they do not represent an improvement on the same complete cohort. No BHS grade is assigned.')
fig,axs=plt.subplots(1,3,figsize=(14,6))
methods=['Relaxed','Strict','B0 constant','B0 on strict subset'];labels=['Relaxed\n17/17','Strict\n11/17','B0 constant\n17/17','B0 constant\nstrict11/17']
for ax,k,upper,lower in zip(axs[:2],['sbp','dbp'],[10,7],[8,5]):
 vals=[next(r['mae'] for r in bp if r['cohort']=='all_retained' and r['method']==m and r['metric']==k) for m in methods];bars=ax.bar(range(4),vals,color=[blue,orange,gray,green]);ax.bar_label(bars,fmt='%.2f',padding=3);ax.axhline(upper,color=red,ls='--',label=f'Upper target {upper}');ax.axhline(lower,color=red,ls=':',label=f'Stronger target {lower}');ax.set_xticks(range(4),labels,rotation=18,ha='right',fontsize=8);ax.set_ylabel('MAE (mmHg)');ax.set_title(k.upper());ax.legend(fontsize=8);ax.set_ylim(0,max(vals)+3)
axs[2].bar(['Relaxed','Strict'],[100,100*11/17],color=[blue,orange]);axs[2].set(ylim=(0,110),ylabel='Evaluation-window coverage (%)',title='Coverage must accompany error');axs[2].text(0,102,'17/17',ha='center');axs[2].text(1,67,'11/17',ha='center')
finish(fig,'F06_BP_gates','F06 | Quality-gate trade-off and the fixed-baseline comparator','Same frozen model; selected quality blocks differ. On the 11 accepted windows, strict and relaxed predictions are identical. A constant B0 comparator outperforms VeinNet even on that same subset. It is a benchmark, not a current BP measurement. Integrity/model/baseline gates are retained in relaxed mode.')
fig,axs=plt.subplots(2,1,figsize=(13,7),sharex=True);xx=np.arange(18)
axs[0].plot(xx,[C[s]['hr'] for s in captured],'-o',color=blue,label='Cuff pulse');axs[0].plot(xx,[P[s]['hr'] for s in captured],'-o',color=orange,label='Tiered IR');axs[0].set_ylabel('HR (BPM)');axs[0].legend()
errors=[P[s]['hr']-C[s]['hr'] for s in captured];axs[1].bar(xx,errors,color=[green if P[s]['hr_tier']=='HIGH' else orange for s in captured]);axs[1].axhline(0,color=gray)
for v in [-10,-5,5,10]:axs[1].axhline(v,color=red,ls='--' if abs(v)==10 else ':',alpha=.6)
axs[1].set_ylabel('Estimate - cuff (BPM)');slot_axis(axs[1],captured)
finish(fig,'F07_HR_trajectory','F07 | HR agreement across rest, motion and challenge slots','All18 recordings, including B0: MAE3.38 BPM, 15/18 within5 BPM and17/18 within10 BPM. Green error bars indicate HIGH; orange ARTIFACT. BPUP2 is the late-reference outlier. The attachment does not define its85% agreement tolerance; both bands are shown, not chosen to manufacture a pass.')
algs=['tiered_IR','MSPTD_ir','Quotient_ir','Elgendi_ir','Welch_ir','Welch3seg_ir','Autocorrelation_ir','MSPTD_red','Quotient_red','Elgendi_red','Welch_red','Welch3seg_red','Autocorrelation_red']
fig,ax=plt.subplots(figsize=(13,8));yy=np.arange(len(algs));rows=old['tables']['HR_comparison']
for gate,off,col in [('off',-.18,blue),('HIGH',.18,orange)]:
 rr=[next(r for r in rows if r['cohort']=='all_captured' and r['algorithm']==name and r['gate']==gate) for name in algs];bars=ax.barh(yy+off,[r['mae'] for r in rr],height=.33,color=col,label='No quality mask' if gate=='off' else 'Shared HIGH mask')
 for i,r in enumerate(rr):ax.text(r['mae']+.12,i+off,f"{r['mae']:.2f} (n={r['n']})",fontsize=8,va='center')
ax.set_yticks(yy,algs);ax.invert_yaxis();ax.set_xlim(0,18);ax.set_xlabel('HR MAE (BPM)');ax.axvline(10,color=red,ls='--',label='Attachment upper MAE target');ax.legend(loc='lower right',fontsize=9)
finish(fig,'F08_HR_algorithms','F08 | Six HR methods on both channels, plus the tiered ensemble','Exploratory comparison on the same retained data. HIGH is the fixed IR-ensemble mask shared across algorithms, not a validated RED-specific gate. Gated n=14; ungated n=18 except Elgendi RED n=17. Production strict HR has0/18 outputs and undefined error. No model was trained or algorithm selected for deployment using these labels.')
ct=J['metrics']['change_tracking'];fig,axs=plt.subplots(1,3,figsize=(13,6))
for ax,key,unit in zip(axs,['sbp','dbp','hr'],['mmHg','mmHg','BPM']):
 x=np.arange(3);ax.bar(x-.18,[r['cuff_d'+key] for r in ct],.36,color=blue,label='Cuff change');ax.bar(x+.18,[r['model_d'+key] for r in ct],.36,color=orange,label='Model/HR change');ax.axhline(0,color=gray);ax.set_xticks(x,['BPUP1*\nvs M06','BPREC1\nvs M06','BPUP2**\nvs M12']);ax.set_title(key.upper());ax.set_ylabel('Change ('+unit+')');ax.legend(fontsize=8)
finish(fig,'F09_change_tracking','F09 | Direction and magnitude of attempted BP changes','*BPUP1 reference averages disagreeing readings; **BPUP2 cuff is late/single. PRE is a prior rest slot, not an immediately preceding synchronized reading. BPREC1 uses M06 by protocol. BPREC2 is missing. The attempted maneuvers do not establish controlled or standardized BP increases.')
fig,axs=plt.subplots(1,2,figsize=(13,6));lat=[live[s]['latency_s'] for s in captured if s in live];ss=[s for s in captured if s in live];durs=[(ts(live[s]['acq_end'])-ts(live[s]['acq_start'])).total_seconds() for s in ss]
axs[0].bar(range(17),durs,color=blue);slot_axis(axs[0],ss);axs[0].set(ylabel='Acquisition span (s)',ylim=(0,35),title='Measured edge acquisition span')
axs[1].bar(range(17),lat,color=green);slot_axis(axs[1],ss);axs[1].axhline(np.percentile(lat,95),color=red,ls='--',label=f'p95={np.percentile(lat,95):.3f}s');axs[1].set(ylabel='Post-acquisition emission latency (s)',ylim=(0,.36),title='Driver live log: n=17');axs[1].legend(fontsize=9)
finish(fig,'F10_latency','F10 | Acquisition and processing latency are different measurements','Post-acquisition latency is edge emitted_at minus edge acquisition_end: p50/p95/max0.212/0.274/0.280s. BPUP2 has no live latency row. The attachment target30-120s includes processing/display; command-to-render and UI rendering were not measured, so this figure cannot certify that endpoint requirement.')
fig,axs=plt.subplots(1,2,figsize=(12,6),gridspec_kw={'width_ratios':[2.2,1]});vals=[P[s]['spo2'] for s in captured];axs[0].plot(range(18),vals,'-o',color=gray);slot_axis(axs[0],captured);axs[0].set_ylabel('Uncalibrated SpO2 formula output (%)');axs[0].set_ylim(55,105);axs[1].hist(vals,bins=[60,65,70,75,80,85,90,95,100],color=gray,edgecolor='white');axs[1].set(xlabel='Formula output (%)',ylabel='Windows');axs[1].yaxis.set_major_locator(MaxNLocator(integer=True))
finish(fig,'F11_SpO2_unvalidated','F11 | Experimental oxygen output - accuracy NOT EVALUABLE','UNCALIBRATED_MAXIM_LINEAR: 18/18 captured outputs, range64.8-95.9%, median90.0%. No independent oxygen reference exists. This is a formula-output distribution, not evidence of hypoxemia, accuracy, or fulfillment of the attached +/-3/4-point requirements. No normal-range band is implied.')
for half,ss in enumerate([captured[:9],captured[9:]]):
 fig,axs=plt.subplots(3,3,figsize=(13,11))
 for ax,s in zip(axs.flat,ss):
  for c,col in [('red',red),('ir',blue)]:
   x=raw[s+'_'+c];ax.plot(np.arange(len(x))/100,100*(x/np.median(x)-1),color=col,lw=.65,label=c.upper(),alpha=.85)
  ax.set_title(s+' | '+P[s]['hr_tier'],fontsize=10);ax.set_xlabel('Time (s)',fontsize=8);ax.set_ylabel('Deviation from median DC (%)',fontsize=8);ax.tick_params(labelsize=8)
 axs.flat[0].legend(fontsize=8)
 finish(fig,'F12'+('a' if half==0 else 'b')+'_all_raw_windows',f'F12{chr(97+half)} | Every retained optical window ({half*9+1}-{half*9+9})','Raw RED/IR samples, normalized only for display as100*(sample/median-1); no smoothing, pulse-band filtering, or waveform selection. Each panel uses its own y-scale, so compare morphology within panels, not apparent height across panels. This visual check is not an independent pulse/SQI validation. All18 windows are retained.')
fig,axs=plt.subplots(1,2,figsize=(12,6));spread={r['slot']:r['beat_spread_bpm'] for r in J['per_slot'] if r.get('status')=='OK'}
for ax,k in zip(axs,['sbp','dbp']):
 x=[spread[s] for s in test];y=[abs(P[s][k]-C[s][k]) for s in test];points(ax,test,x,y);rho=spearmanr(x,y).statistic;ax.set(xlabel='Range of six IR HR estimates (BPM)',ylabel='Absolute '+k.upper()+' error (mmHg)',title=f'{k.upper()} | tie-aware Spearman {rho:.3f}')
finish(fig,'F13_quality_vs_BP','F13 | Exploratory estimator disagreement versus BP error','n=17 post-baseline windows; estimator ranges rounded to 0.01 BPM as in the original analyzer. Correlation is descriptive with ties handled correctly; no BP acceptance threshold was fitted. Good HR agreement does not by itself guarantee accurate BP. This addresses the objective linking signal/HR quality with BP while preserving their distinct validation requirements.')
atlas.close()
span=(ts(sealed['BPUP2']['ppg_end_utc'])-start).total_seconds()
summary=dict(n_captured=18,n_planned=19,n_bp=17,raw_pairs=54000,sampled_s=540,first_local=local(sealed['B0']['ppg_start_utc']),last_local=local(sealed['BPUP2']['ppg_end_utc']),span_s=span,calibration_cuff=C['B0'],bp=bp,hr=hr,coverage=coverage,figures=figures,requirements=req,tables={k:len(v) for k,v in tables.items()},latency=dict(n=len(lat),p50=float(np.percentile(lat,50)),p95=float(np.percentile(lat,95)),max=max(lat)))
(O/'results_package.json').write_text(json.dumps(summary,indent=2,default=lambda x:float(x)))
(O/'FIGURE_CAPTIONS.md').write_text('# FAST1 figure captions\n\n'+'\n\n'.join('## '+f['title']+'\n\n'+f['caption'] for f in figures))
# Verify spreadsheet is readable and required data have not been dropped.
v=load_workbook(O/'FAST1_RESULTS_TABLES.xlsx',read_only=True,data_only=True);assert v['T04_all_slots'].max_row==20 and v['T11_cuff_readings'].max_row==21;v.close()
assert abs(next(r['mae'] for r in bp if r['cohort']=='all_retained' and r['method']=='Relaxed' and r['metric']=='sbp')-13.363463)<1e-5
print('Created',len(figures),'figures (PNG/SVG/PDF),',len(tables),'tables, atlas and workbook; span_s=',span)
