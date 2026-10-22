#!/usr/bin/env python3
"""FAST1 offline analysis: raw journals + cuff events -> FAST1_RESULTS.xlsx, FAST1_RESULTS_HE.md, fast1_results.json.

Replays every slot window through the REAL edge pipeline (same code path as prototype_continuous_20261020/replay_all.py),
once in prototype-1.1.0 mode and once strict (flag unset). Cuff values enter ONLY the baseline calibration and the error
columns, never a target prediction. Needs the runtime venv (scipy/onnx/openpyxl).

  python analyze_fast1.py --run ~/VeinnovateRuns/POC-20261021-FAST1 --project <repo> --out <dir>
Legacy runs: add --plan legacy --baseline B0R --cuff-json cuffs.json
"""
import argparse, asyncio, collections, datetime as dt, hashlib, json, os, sys, tempfile, time
from pathlib import Path
import numpy as np

HERE = Path(__file__).resolve().parent
UID, AGE, SEX = "A90AFDAF672D66F1", 32, "M"   # SUBJ-001 (same as replay_all)
CAL = ("B0", "B0R", "B0R2")


# ------------------------------------------------------------------ inputs
def load_plan(spec):
    if spec == "legacy":
        return dict(slots=[], baseline_slots=list(CAL), p95_latency_limit_s=2.0)
    return json.loads(Path(spec).read_text())


def read_cuffs(run, override=None):
    """slot -> dict(sbp,dbp,hr,n). Last SAVE event wins; 'readings' (AVG2) are averaged, else 'values'."""
    out = {}
    p = Path(run) / "cuff_events.jsonl"
    for line in (p.read_text().splitlines() if p.exists() else []):
        e = json.loads(line)
        if e.get("kind") != "SAVE":
            continue
        rd = e.get("readings") or [e["values"]]
        out[e["slot"]] = dict(sbp=float(np.mean([r[0] for r in rd])), dbp=float(np.mean([r[1] for r in rd])),
                              hr=float(np.mean([r[2] for r in rd])) if all(len(r) > 2 and r[2] is not None for r in rd) else None, n=len(rd))
    for s, v in (override or {}).items():
        out[s] = dict(sbp=v[0], dbp=v[1], hr=v[2], n=1)
    return out


def read_raw(run, run_id=None):
    """slot -> DATA records (hash-verified against raw_manifest.json); slot_map.jsonl (driver) fills untagged records."""
    run = Path(run)
    recs = []
    for e in json.loads((run / "raw_manifest.json").read_text())["files"]:
        b = (run / e["path"]).read_bytes()
        assert hashlib.sha256(b).hexdigest() == e["sha256"], f"raw hash mismatch {e['path']}"
        recs += [json.loads(l) for l in b.splitlines()]
    if run_id:   # never mix runs: every journal record must belong to the run being analysed (BENCH data can never enter FAST1)
        foreign = {r.get("run_id") for r in recs} - {run_id}
        if foreign: raise ValueError(f"raw journal contains records of other run(s) {sorted(map(str, foreign))}; refusing to analyse as {run_id}")
    smap = [json.loads(l) for l in (run / "slot_map.jsonl").read_text().splitlines()] if (run / "slot_map.jsonl").exists() else []
    groups = collections.defaultdict(list)
    for r in recs:
        if r.get("channel") != "DATA" or not r.get("packet_hex"):
            continue
        s = r.get("slot")
        if not s or s == "NONE":
            s = next((m["slot"] for m in smap if m["start_utc"] <= r["t_utc"] <= m["end_utc"]), None)
        if s:
            groups[s].append(r)
    return groups


# ------------------------------------------------------------------ pipeline (mirrors replay_all.py)
def build_windows(groups):
    from protocol_decoder import decode_data_packet
    from window_assembler import SessionRingBuffer
    W = []
    for slot, data in groups.items():
        ring = SessionRingBuffer()
        for r in data:
            ring.ingest(decode_data_packet(bytes.fromhex(r["packet_hex"])))
        w = ring.try_cut_window(100)
        if w is None or not w.integrity.is_complete:
            continue
        W.append(dict(slot=slot, t0=data[0]["t_utc"], red=np.array(w.red, float), ir=np.array(w.ir, float)))
    return sorted(W, key=lambda w: w["t0"])


def replay(W, cuffs, baseline_slots, mode):
    import prototype_continuous as proto
    from calibration_store import CalibrationStore
    from veinnet_edge_service import EdgeOrchestrator
    from veinnet_sibling_adapter import VeinNetSiblingAdapter
    from window_assembler import AssembledWindow, WindowIntegrity
    def mkwin(red, ir):
        return AssembledWindow(sensor_session_id=1, sensor_config_id=25, device_id=int(UID, 16), first_sample_index=0, last_sample_index=2999,
            red=[int(round(v)) for v in red[:3000]], ir=[int(round(v)) for v in ir[:3000]], sample_indices=list(range(3000)),
            integrity=WindowIntegrity(expected_samples=3000, received_samples=3000, duplicate_samples=0, lost_samples=0, out_of_order_packets=0,
                session_changed=False, config_changed=False, fifo_overflow_count=0, loss_evidence_available=True))
    adapter = VeinNetSiblingAdapter()
    if mode == "proto":
        os.environ[proto.ENV_FLAG] = "1"
    else:
        os.environ.pop(proto.ENV_FLAG, None)
    o = EdgeOrchestrator(model_adapter=adapter, calibration_store=CalibrationStore(Path(tempfile.mkdtemp()) / "cal.json"),
                         patient_id="SUBJ-001", expected_hardware_uid=UID, short_code="VN-001", logical_session_id="FAST1")
    out = []
    async def sink(r): out.append(r)
    o.result_listeners.append(sink)
    base = next((w for w in W if w["slot"] in baseline_slots and w["slot"] in cuffs), None)   # first baseline window is the calibration
    cal_note, rows = "no BP calibration", {}
    for w in W:
        if w is base:
            bw = mkwin(w["red"], w["ir"])
            pre = adapter.preprocess_raw_window(bw, role="baseline") if mode == "proto" else adapter.preprocess_raw_window(bw)
            if pre.accepted:
                c = cuffs[w["slot"]]
                o.calibration_store.capture(patient_id=o.patient_id, device_id=UID, sensor_session_id=1, sensor_config_id=25,
                    baseline_cuff_sbp=c["sbp"], baseline_cuff_dbp=c["dbp"], baseline_ppg_red=bw.red, baseline_ppg_ir=bw.ir,
                    baseline_ppg_processed=pre.ppg_flat, baseline_morphology=pre.morphology, age_years=AGE, sex=SEX,
                    firmware_version="3.0.1", preprocessing_version=adapter.preprocessing_version)
                cal_note = f"calibrated on {w['slot']}"
            else:
                cal_note = f"{w['slot']} REJECTED: {pre.reason_code}"
        n0 = len(out); t0 = time.perf_counter()
        asyncio.run(o._process_window(mkwin(w["red"], w["ir"])))
        ms = (time.perf_counter() - t0) * 1000.0
        r = out[-1].to_dict() if len(out) > n0 else None
        ht = getattr(o, "last_prototype_hr", None) if mode == "proto" else None
        rows[w["slot"]] = dict(bp_state=r["result_state"] if r else None, sbp=r["bp"]["sbp_mmhg"] if r else None, dbp=r["bp"]["dbp_mmhg"] if r else None,
            warnings=";".join(r["sqi"]["reason_codes"]) if r else None, hr=r["hr"]["bpm"] if r else None, hr_valid=r["hr"]["valid"] if r else None,
            hr_tier=ht.tier if ht else None, hr_components=ht.components if ht else None,
            spo2=r["spo2"]["pct"] if r else None, spo2_status=r["spo2"]["status"] if r else None,
            sqi_version=r["provenance"]["sqi_version"] if r else None, model_sha=r["provenance"]["model_sha256"][:8] if r else None,
            latency_ms=round(ms, 1), calibration=cal_note)
        if mode == "proto":
            o.last_prototype_hr = None; o.last_prototype_spo2 = None
    os.environ.pop(proto.ENV_FLAG, None)
    return rows, adapter.model_sha256, proto.VERSION


# ------------------------------------------------------------------ metrics
def has(x): return x is not None


def err_stats(pairs):
    """pairs of (model, ref) -> n, MAE, bias, SD, % within 5/10/15."""
    if not pairs:
        return dict(n=0)
    e = np.array([m - r for m, r in pairs], float); a = np.abs(e)
    return dict(n=len(e), mae=round(float(a.mean()), 2), bias=round(float(e.mean()), 2), sd=round(float(e.std(ddof=1)), 2) if len(e) > 1 else None,
                within5=round(100 * float((a <= 5).mean()), 1), within10=round(100 * float((a <= 10).mean()), 1), within15=round(100 * float((a <= 15).mean()), 1))


def pct(a, b): return f"{a}/{b}"


def rank_corr(x, y):
    if len(x) < 3:
        return None
    rx, ry = np.argsort(np.argsort(x)), np.argsort(np.argsort(y))
    return round(float(np.corrcoef(rx, ry)[0, 1]), 3) if rx.std() and ry.std() else None


def bland_altman(pairs):
    if len(pairs) < 2:
        return dict(n=len(pairs))
    d = np.array([m - r for m, r in pairs], float); sd = float(d.std(ddof=1))
    return dict(n=len(d), bias=round(float(d.mean()), 2), sd=round(sd, 2), loa_low=round(float(d.mean() - 1.96 * sd), 2), loa_high=round(float(d.mean() + 1.96 * sd), 2),
                points=[(round((m + r) / 2, 2), round(m - r, 2)) for m, r in pairs])


def live_metrics(run, plan):
    """Continuity/latency from the driver's own logs (absent for legacy runs -> None)."""
    run = Path(run); lp, ep = run / "live_results.jsonl", run / "events.jsonl"
    if not lp.exists():
        return None
    L = [json.loads(l) for l in lp.read_text().splitlines() if l]
    ev = [json.loads(l) for l in ep.read_text().splitlines() if l] if ep.exists() else []
    lat = np.array([r["latency_s"] for r in L if has(r.get("latency_s"))], float)
    out = dict(windows_seen=len(L), expected=len(plan["slots"]), gaps=sum(e["kind"] == "GAP" for e in ev), ble_drops=sum(e["kind"] == "BLE_DROP" for e in ev),
               exceptions=sum(e["kind"] == "EXCEPTION" for e in ev))
    if len(lat):
        out |= dict(p50_s=round(float(np.percentile(lat, 50)), 3), p95_s=round(float(np.percentile(lat, 95)), 3), max_s=round(float(lat.max()), 3),
                    p95_pass=bool(np.percentile(lat, 95) <= plan["p95_latency_limit_s"]))
    return out


def analyze(run, plan, cuffs, baseline_slots, W=None, proto_rows=None, strict_rows=None):
    plan_slots = [s["slot"] for s in plan["slots"]]
    slot_info = {s["slot"]: s for s in plan["slots"]}
    order = plan_slots + [w["slot"] for w in W if w["slot"] not in plan_slots]
    P, S = proto_rows, strict_rows
    scored = [s for s in order if s in P and s not in CAL and s in cuffs]              # comparison slots with a cuff
    base = next((s for s in baseline_slots if s in cuffs and s in P), None)
    bp = [s for s in scored if P[s]["bp_state"] == "VALID" and has(P[s]["sbp"])]
    cf = cuffs.get(base) if base else None
    M = {}
    M["bp_main"] = dict(sbp=err_stats([(P[s]["sbp"], cuffs[s]["sbp"]) for s in bp]), dbp=err_stats([(P[s]["dbp"], cuffs[s]["dbp"]) for s in bp]),
                        bp_on_cuff_slots=pct(len(bp), len(scored)))
    for k in ("sbp", "dbp"):
        m = M["bp_main"][k]
        m["target_mae_le_5"] = ("MET" if m["mae"] <= 5 else "NOT MET") if m["n"] else "n/a"
    M["bp_carry_forward_from_baseline"] = dict(baseline=base, sbp=err_stats([(cf["sbp"], cuffs[s]["sbp"]) for s in bp]) if cf else {},
                                               dbp=err_stats([(cf["dbp"], cuffs[s]["dbp"]) for s in bp]) if cf else {})
    # change tracking: slot vs its PRE slot
    ct = []
    for s in plan_slots:
        pre = slot_info[s].get("pre")
        if pre and s in P and pre in P and s in cuffs and pre in cuffs and all(has(P[x]["sbp"]) for x in (s, pre)):
            dc = (cuffs[s]["sbp"] - cuffs[pre]["sbp"], cuffs[s]["dbp"] - cuffs[pre]["dbp"])
            dm = (P[s]["sbp"] - P[pre]["sbp"], P[s]["dbp"] - P[pre]["dbp"])
            ct.append(dict(slot=s, pre=pre, cuff_dsbp=round(dc[0], 1), model_dsbp=round(dm[0], 1), abs_err_sbp=round(abs(dm[0] - dc[0]), 1), carry_abs_err_sbp=round(abs(dc[0]), 1),
                           cuff_ddbp=round(dc[1], 1), model_ddbp=round(dm[1], 1), abs_err_dbp=round(abs(dm[1] - dc[1]), 1), carry_abs_err_dbp=round(abs(dc[1]), 1),
                           cuff_dhr=round(cuffs[s]["hr"] - cuffs[pre]["hr"], 1) if has(cuffs[s].get("hr")) and has(cuffs[pre].get("hr")) else None,
                           model_dhr=round(P[s]["hr"] - P[pre]["hr"], 1) if has(P[s]["hr"]) and has(P[pre]["hr"]) else None))
    M["change_tracking"] = ct
    # HR
    hc = [s for s in order if s in P and s in cuffs and has(cuffs[s].get("hr")) and has(P[s]["hr"])]
    M["hr"] = dict(all=err_stats([(P[s]["hr"], cuffs[s]["hr"]) for s in hc]),
                   by_tier={t: err_stats([(P[s]["hr"], cuffs[s]["hr"]) for s in hc if P[s]["hr_tier"] == t]) for t in ("HIGH", "LOW", "ARTIFACT")},
                   tier_counts={t: sum(P[s]["hr_tier"] == t for s in P) for t in ("HIGH", "LOW", "ARTIFACT")})
    sp = [P[s]["spo2"] for s in P if has(P[s]["spo2"])]
    M["spo2"] = dict(displayed=pct(len(sp), len(P)), label="UNCALIBRATED_MAXIM_LINEAR", accuracy="NOT_EVALUABLE",
                     min=min(sp) if sp else None, median=float(np.median(sp)) if sp else None, max=max(sp) if sp else None)
    M["continuity"] = dict(windows_with_data=len(P), expected_slots=len(plan_slots) or len(P),
                           pct_bp=round(100 * sum(P[s]["bp_state"] == "VALID" for s in P) / max(len(P), 1), 1),
                           pct_hr=round(100 * sum(bool(P[s]["hr_valid"]) for s in P) / max(len(P), 1), 1), pct_spo2=round(100 * len(sp) / max(len(P), 1), 1),
                           missing_slots=[s for s in plan_slots if s not in P], live=live_metrics(run, plan))
    lat = np.array([P[s]["latency_ms"] for s in P], float)
    M["latency_offline_processing_ms"] = dict(p50=round(float(np.percentile(lat, 50)), 1), p95=round(float(np.percentile(lat, 95)), 1), max=round(float(lat.max()), 1)) if len(lat) else {}
    # offline strict like-for-like: windows strict answered
    ll_bp = [s for s in bp if S[s]["bp_state"] == "VALID" and has(S[s]["sbp"])]
    ll_hr = [s for s in hc if has(S[s]["hr"])]
    M["like_for_like_strict"] = dict(
        bp_n=len(ll_bp), bp_strict_sbp=err_stats([(S[s]["sbp"], cuffs[s]["sbp"]) for s in ll_bp]), bp_proto_sbp=err_stats([(P[s]["sbp"], cuffs[s]["sbp"]) for s in ll_bp]),
        bp_strict_dbp=err_stats([(S[s]["dbp"], cuffs[s]["dbp"]) for s in ll_bp]), bp_proto_dbp=err_stats([(P[s]["dbp"], cuffs[s]["dbp"]) for s in ll_bp]),
        hr_n=len(ll_hr), hr_strict=err_stats([(S[s]["hr"], cuffs[s]["hr"]) for s in ll_hr]), hr_proto=err_stats([(P[s]["hr"], cuffs[s]["hr"]) for s in ll_hr]),
        strict_answered_bp=pct(sum(S[s]["bp_state"] == "VALID" for s in scored), len(scored)), strict_answered_hr=pct(sum(bool(S[s]["hr_valid"]) for s in S), len(S)))
    # beat consistency = range (BPM) of the six IR HR estimators; lower = more consistent
    def spread(s):
        c = P[s].get("hr_components")
        v = [x for x in (c or {}).values() if isinstance(x, (int, float))]
        return round(max(v) - min(v), 2) if v else None
    pts = [(spread(s), abs(P[s]["sbp"] - cuffs[s]["sbp"]), abs(P[s]["dbp"] - cuffs[s]["dbp"])) for s in bp if has(spread(s))]
    M["beat_consistency"] = dict(n=len(pts), definition="range of six IR HR estimators (BPM); lower = more consistent",
                                 spearman_vs_abs_sbp_err=rank_corr([p[0] for p in pts], [p[1] for p in pts]),
                                 spearman_vs_abs_dbp_err=rank_corr([p[0] for p in pts], [p[2] for p in pts]),
                                 note="FAST1 data only; candidate BP quality label, not a threshold")
    M["bland_altman"] = dict(sbp=bland_altman([(P[s]["sbp"], cuffs[s]["sbp"]) for s in bp]), dbp=bland_altman([(P[s]["dbp"], cuffs[s]["dbp"]) for s in bp]),
                             hr=bland_altman([(P[s]["hr"], cuffs[s]["hr"]) for s in hc]))
    rows = []
    for s in order:
        if s not in P:
            rows.append(dict(slot=s, kind=slot_info.get(s, {}).get("kind"), status="NO_DATA"))
            continue
        p, c = P[s], cuffs.get(s, {})
        rows.append(dict(slot=s, kind=slot_info.get(s, {}).get("kind"), status="OK", window_start_utc=next((w["t0"] for w in W if w["slot"] == s), None),
                         cuff_sbp=c.get("sbp"), cuff_dbp=c.get("dbp"), cuff_hr=c.get("hr"), cuff_n=c.get("n"),
                         model_sbp=p["sbp"], model_dbp=p["dbp"], bp_state=p["bp_state"], warnings=p["warnings"], hr=p["hr"], hr_tier=p["hr_tier"],
                         spo2=p["spo2"], spo2_label="UNCALIBRATED_MAXIM_LINEAR" if has(p["spo2"]) else None, beat_spread_bpm=spread(s),
                         strict_sbp=S[s]["sbp"], strict_dbp=S[s]["dbp"], strict_hr=S[s]["hr"], latency_ms=p["latency_ms"], calibration=p["calibration"]))
    return M, rows


# ------------------------------------------------------------------ outputs
def write_xlsx(path, M, rows, meta):
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    wb = Workbook(); ws = wb.active; ws.title = "PerSlot"
    cols = list(dict.fromkeys(k for r in rows for k in r)); ws.append(cols)
    for r in rows:
        ws.append([r.get(c) for c in cols])
    sm = wb.create_sheet("Summary")
    def flat(prefix, d):
        for k, v in d.items():
            if isinstance(v, dict) and k != "points":
                yield from flat(f"{prefix}{k}.", v)
            elif k != "points" and not isinstance(v, list):
                yield prefix + k, v
    for k, v in flat("", {m: M[m] for m in M if m not in ("change_tracking", "bland_altman")}):
        sm.append([k, v])
    for k, v in meta.items():
        sm.append(["meta." + k, v])
    ct = wb.create_sheet("ChangeTracking")
    if M["change_tracking"]:
        ct.append(list(M["change_tracking"][0]))
        for r in M["change_tracking"]:
            ct.append(list(r.values()))
    for name, key in (("BA_SBP", "sbp"), ("BA_DBP", "dbp"), ("BA_HR", "hr")):
        b = M["bland_altman"][key]; w = wb.create_sheet(name)
        w.append(["mean(model,cuff)", "model-cuff"])
        for x, y in b.get("points", []):
            w.append([x, y])
        w.append([]); w.append(["bias", b.get("bias")]); w.append(["loa_low", b.get("loa_low")]); w.append(["loa_high", b.get("loa_high")])
        n = len(b.get("points", []))
        if n:
            ch = ScatterChart(); ch.title = f"Bland-Altman {key.upper()} (model - cuff)"; ch.style = 13
            ch.x_axis.title = "mean of model and cuff"; ch.y_axis.title = "model - cuff"
            s = Series(Reference(w, min_col=2, min_row=2, max_row=n + 1), Reference(w, min_col=1, min_row=2, max_row=n + 1))
            s.marker.symbol = "circle"; s.graphicalProperties.line.noFill = True; ch.series.append(s)
            w.add_chart(ch, "D2")
    wb.save(path)


def write_he(path, M, meta):
    b, c, h, sp, ct = M["bp_main"], M["continuity"], M["hr"], M["spo2"], M["change_tracking"]
    f = lambda d, k: d.get(k, "—") if d.get("n") else "—"
    L = [f"# FAST1 — סיכום תוצאות ({meta['run_id']})", "",
         f"גרסה `{meta['prototype_version']}` · מודל `{meta['model_sha256'][:8]}` · SpO2: מוצג תמיד, `UNCALIBRATED_MAXIM_LINEAR`, ללא אוקסימטר — דיוק `NOT_EVALUABLE`.", "",
         "## לחץ דם (BP)",
         f"- SBP: MAE {f(b['sbp'],'mae')} · הטיה {f(b['sbp'],'bias')} · SD {f(b['sbp'],'sd')} · בטווח 5/10/15: {f(b['sbp'],'within5')}/{f(b['sbp'],'within10')}/{f(b['sbp'],'within15')}% · יעד ≤5: **{b['sbp']['target_mae_le_5']}** (n={b['sbp']['n']})",
         f"- DBP: MAE {f(b['dbp'],'mae')} · הטיה {f(b['dbp'],'bias')} · SD {f(b['dbp'],'sd')} · בטווח 5/10/15: {f(b['dbp'],'within5')}/{f(b['dbp'],'within10')}/{f(b['dbp'],'within15')}% · יעד ≤5: **{b['dbp']['target_mae_le_5']}** (n={b['dbp']['n']})",
         f"- נקודת ייחוס (carry-forward מ-B0): SBP MAE {f(M['bp_carry_forward_from_baseline']['sbp'],'mae')} · DBP MAE {f(M['bp_carry_forward_from_baseline']['dbp'],'mae')}",
         f"- חלונות עם BP: {b['bp_on_cuff_slots']}", "", "## מעקב שינוי (BPUP/BPREC מול PRE)"]
    L += [f"- {r['slot']} מול {r['pre']}: ΔSBP כרית {r['cuff_dsbp']} / מודל {r['model_dsbp']} (|שגיאה| {r['abs_err_sbp']}, carry {r['carry_abs_err_sbp']}); ΔDBP כרית {r['cuff_ddbp']} / מודל {r['model_ddbp']} (|שגיאה| {r['abs_err_dbp']}); ΔHR {r['cuff_dhr']} / {r['model_dhr']}" for r in ct] or ["- אין נתונים"]
    L += ["", "## דופק (HR)", f"- MAE מול כרית: {f(h['all'],'mae')} · בטווח ±5: {f(h['all'],'within5')}% (n={h['all']['n']})"]
    L += [f"- דרגה {t}: MAE {f(v,'mae')} (n={v['n']}) · מספר חלונות {h['tier_counts'][t]}" for t, v in h["by_tier"].items()]
    L += ["", "## SpO2 (ניסיוני בלבד)", f"- מוצג ב-{sp['displayed']} חלונות · טווח {sp['min']}–{sp['max']} · חציון {sp['median']}", "",
          "## רציפות והשהיה", f"- חלונות עם נתונים {c['windows_with_data']}/{c['expected_slots']} · BP {c['pct_bp']}% · HR {c['pct_hr']}% · SpO2 {c['pct_spo2']}%",
          f"- חסרים: {', '.join(c['missing_slots']) or 'אין'}"]
    lv = c["live"]
    L += [f"- חי: פער {lv['gaps']} · ניתוקי BLE {lv['ble_drops']} · חריגות {lv['exceptions']} · p50/p95/max {lv.get('p50_s')}/{lv.get('p95_s')}/{lv.get('max_s')} שנ' (גבול p95 ≤ 2)" if lv else
          "- אין לוג חי (ריצה היסטורית) — רק זמן עיבוד offline: " + json.dumps(M["latency_offline_processing_ms"])]
    ll = M["like_for_like_strict"]
    L += ["", "## השוואה הוגנת מול strict (אותם חלונות)", f"- BP: strict ענה {ll['strict_answered_bp']}; על {ll['bp_n']} משותפים SBP MAE strict {f(ll['bp_strict_sbp'],'mae')} / 1.1.0 {f(ll['bp_proto_sbp'],'mae')}",
          f"- HR: על {ll['hr_n']} משותפים MAE strict {f(ll['hr_strict'],'mae')} / 1.1.0 {f(ll['hr_proto'],'mae')}", "",
          "## עקביות פעימות מול שגיאת BP", f"- n={M['beat_consistency']['n']} · Spearman SBP {M['beat_consistency']['spearman_vs_abs_sbp_err']} · DBP {M['beat_consistency']['spearman_vs_abs_dbp_err']} — {M['beat_consistency']['note']}", "",
          "## Bland-Altman"] + [f"- {k.upper()}: הטיה {v.get('bias')} · גבולות הסכמה {v.get('loa_low')} … {v.get('loa_high')} (n={v['n']})" for k, v in M["bland_altman"].items()]
    Path(path).write_text("\n".join(L) + "\n", encoding="utf-8")


def live_vs_replay(run, P, bp_tol=1.0):
    """Live readings (live_results.jsonl, from the edge service) vs the offline replay of the same raw windows.
    HR value, HR tier and SpO2 must be identical; BP within bp_tol mmHg (replay assumes the SUBJ-001 profile 32/M; a larger BP gap is
    reported as a profile/calibration difference, not hidden). Returns dict(ok, detail, rows)."""
    live = {}
    p = Path(run) / "live_results.jsonl"
    for l in (p.read_text().splitlines() if p.exists() else []):
        d = json.loads(l); live[d["slot"]] = d
    rows, bad = [], []
    for slot, d in live.items():
        r = P.get(slot)
        if r is None:
            bad.append(f"{slot}: no replay"); continue
        same_hr = abs((d["hr"] or 0) - (r["hr"] or 0)) < 0.05 and (d["tier"] or "").replace("PROTOTYPE_", "").replace("_CONFIDENCE", "") == (r["hr_tier"] or "")
        same_sp = d["spo2"] is not None and r["spo2"] is not None and abs(d["spo2"] - r["spo2"]) < 0.05
        dbp = max(abs((d[k] or 0) - (r[k] or 0)) for k in ("sbp", "dbp")) if d["sbp"] is not None and r["sbp"] is not None else (0 if d["sbp"] is None and r["sbp"] is None else 99)
        rows.append(dict(slot=slot, live_hr=d["hr"], replay_hr=r["hr"], live_tier=d["tier"], replay_tier=r["hr_tier"], live_spo2=d["spo2"], replay_spo2=r["spo2"],
                         live_sbp=d["sbp"], replay_sbp=r["sbp"], max_bp_gap=round(dbp, 3)))
        if not (same_hr and same_sp and dbp <= bp_tol): bad.append(f"{slot}: hr={same_hr} spo2={same_sp} bp_gap={dbp:.2f}")
    return dict(ok=bool(rows) and not bad, detail=f"{len(rows)} windows compared" + (f"; DIFF {bad}" if bad else "; identical"), rows=rows)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True); ap.add_argument("--project", required=True); ap.add_argument("--bundles")
    ap.add_argument("--plan", default=str(HERE / "fast1_plan.json")); ap.add_argument("--baseline", default=None, help="baseline slot (default: plan baseline_slots order)")
    ap.add_argument("--cuff-json"); ap.add_argument("--out", default=str(HERE))
    a = ap.parse_args(argv)
    B = Path(a.bundles or a.project)
    sys.path[:0] = [str(B / "Hardware_integration_ready" / "receiver"), str(B / "VeinNet_integration_ready")]
    plan = load_plan(a.plan); run = Path(a.run).expanduser()
    cuffs = read_cuffs(run, json.loads(Path(a.cuff_json).read_text()) if a.cuff_json else None)
    if plan.get("run_id") and run.name != plan["run_id"]: ap.error(f"run folder {run.name} != plan run_id {plan['run_id']}")
    W = build_windows(read_raw(run, plan.get("run_id")))
    bl = [a.baseline] if a.baseline else plan["baseline_slots"]
    P, model, ver = replay(W, cuffs, bl, "proto"); S, _, _ = replay(W, cuffs, bl, "strict")
    M, rows = analyze(run, plan, cuffs, bl, W, P, S)
    meta = dict(run_id=run.name, prototype_version=ver, model_sha256=model, host=os.uname().nodename + "/" + os.uname().machine,
                generated_utc=dt.datetime.now(dt.timezone.utc).isoformat(), n_windows=len(W))
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    (out / "fast1_results.json").write_text(json.dumps(dict(meta=meta, metrics=M, per_slot=rows, proto=P, strict=S), indent=1, default=str))
    write_xlsx(out / "FAST1_RESULTS.xlsx", M, rows, meta); write_he(out / "FAST1_RESULTS_HE.md", M, meta)
    print(json.dumps(dict(meta=meta, bp=M["bp_main"], hr=M["hr"]["all"]), default=str)[:600])
    return M, rows, P, S


if __name__ == "__main__":
    main()
